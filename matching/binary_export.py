import io
import os
import zipfile
from datetime import datetime, timezone
from typing import Optional, List
from matching.models import TailoredArtifact, ArtifactType
from matching.templates import Template, get_template
import docx
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# BRIEF-FR-006 council review #2 MAJOR 6: the requirement is that export is
# deterministic (same artifact -> same bytes). A fixed, arbitrary reference
# instant, never wall-clock `datetime.now()`.
_DETERMINISTIC_DT = datetime(2026, 1, 1, tzinfo=timezone.utc)
_DETERMINISTIC_ZIP_DATE_TIME = (2026, 1, 1, 0, 0, 0)


def _make_docx_bytes_deterministic(data: bytes) -> bytes:
    """python-docx's `Document.save()` writes each zip member with the
    current wall-clock time in its `ZipInfo.date_time`, so two exports of the
    identical artifact differ byte-for-byte. Rewrite every member with a
    fixed timestamp, preserving content and compression exactly."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out_bio = io.BytesIO()
    with zipfile.ZipFile(out_bio, "w", zipfile.ZIP_DEFLATED) as out:
        for info in src.infolist():
            info.date_time = _DETERMINISTIC_ZIP_DATE_TIME
            out.writestr(info, src.read(info.filename))
    src.close()
    return out_bio.getvalue()


class BinaryArtifactExporter:
    @staticmethod
    def export_to_docx(
        artifact: TailoredArtifact, output_path: Optional[str] = None, template: str | Template | None = None,
    ) -> bytes:
        tpl = template if isinstance(template, Template) else get_template(template)
        doc = docx.Document()

        # Page Margins: template-controlled for clean ATS presentation.
        # Single column throughout -- no `doc.add_table()` is ever called
        # (no table is used for layout, per every committed template).
        for s in doc.sections:
            s.top_margin = Inches(tpl.margin_in)
            s.bottom_margin = Inches(tpl.margin_in)
            s.left_margin = Inches(tpl.margin_in)
            s.right_margin = Inches(tpl.margin_in)

        # Document Title / Header
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(artifact.title)
        title_run.bold = True
        title_run.font.size = Pt(tpl.title_size_pt)
        title_run.font.name = tpl.docx_font

        for sec in artifact.sections:
            if sec.heading:
                # Real heading style (Word's built-in "Heading 2"), not a
                # bold paragraph styled to merely look like one.
                h_p = doc.add_paragraph(style="Heading 2")
                h_run = h_p.add_run(sec.heading)
                h_run.bold = True
                h_run.font.size = Pt(tpl.heading_size_pt)
                h_run.font.name = tpl.docx_font

            if sec.content:
                p = doc.add_paragraph()
                r = p.add_run(sec.content)
                r.font.size = Pt(tpl.body_size_pt)
                r.font.name = tpl.docx_font

            for item in sec.items:
                bp = doc.add_paragraph(style='List Bullet')
                br = bp.add_run(item)
                br.font.size = Pt(tpl.bullet_size_pt)
                br.font.name = tpl.docx_font

        doc.core_properties.created = _DETERMINISTIC_DT
        doc.core_properties.modified = _DETERMINISTIC_DT
        doc.core_properties.last_modified_by = ""

        bio = io.BytesIO()
        doc.save(bio)
        content = _make_docx_bytes_deterministic(bio.getvalue())

        if output_path:
            with open(output_path, "wb") as f:
                f.write(content)

        return content

    @staticmethod
    def export_to_pdf(
        artifact: TailoredArtifact, output_path: Optional[str] = None, template: str | Template | None = None,
    ) -> bytes:
        tpl = template if isinstance(template, Template) else get_template(template)
        bio = io.BytesIO()
        doc = SimpleDocTemplate(
            bio,
            pagesize=letter,
            rightMargin=tpl.margin_pt,
            leftMargin=tpl.margin_pt,
            topMargin=tpl.margin_pt,
            bottomMargin=tpl.margin_pt,
            # MAJOR 6: suppresses reportlab's per-build CreationDate/ModDate/ID
            # (council-verified to make repeat exports byte-identical).
            invariant=1,
        )
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Heading1'],
            fontSize=tpl.title_size_pt,
            leading=tpl.title_size_pt + 4,
            alignment=TA_LEFT,
            fontName=tpl.pdf_font_bold,
            spaceAfter=12,
        )
        heading_style = ParagraphStyle(
            'SectionHeading',
            parent=styles['Heading2'],
            fontSize=tpl.heading_size_pt,
            leading=tpl.heading_size_pt + 4,
            alignment=TA_LEFT,
            fontName=tpl.pdf_font_bold,
            spaceBefore=10,
            spaceAfter=4,
        )
        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontSize=tpl.body_size_pt,
            leading=tpl.body_size_pt + 4,
            fontName=tpl.pdf_font,
            spaceAfter=6,
        )
        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontSize=tpl.bullet_size_pt,
            leading=tpl.bullet_size_pt + 4,
            fontName=tpl.pdf_font,
            leftIndent=15,
            spaceAfter=3,
        )

        # Single column, no tables or text boxes: `story` is a flat,
        # top-to-bottom flowable list (Paragraph/Spacer only). No
        # `reportlab.platypus.Table` is imported anywhere in this module.
        story = []
        story.append(Paragraph(artifact.title, title_style))

        for sec in artifact.sections:
            if sec.heading:
                story.append(Paragraph(sec.heading, heading_style))
            if sec.content:
                story.append(Paragraph(sec.content.replace("\n", "<br/>"), body_style))
            for item in sec.items:
                story.append(Paragraph(f"- {item}", bullet_style))

        doc.build(story)
        content = bio.getvalue()

        if output_path:
            with open(output_path, "wb") as f:
                f.write(content)

        return content
